"""
Generación de reportes server-side (F3, REPORTS_EXPORT). Lee el ROLLUP vía `metrics.py` (NO escanea
las fuentes — reusa get_summary/get_funnel/get_appointments_distribution/get_leads_evolution, que
ya resuelven labels/colores del catálogo y validan el rango) y arma un binario PDF (reportlab) o
Excel (openpyxl). Devuelve `(content, media_type, filename)`; el ROUTER arma el `Response`.

Es una de las dos excepciones a "los services devuelven el envelope": el binario es transporte, no
payload de dominio (igual que el webhook de conversations devuelve un `Response` crudo).

`reportlab`/`openpyxl` se importan LAZY dentro de cada generador → el boot/smoke/mypy sin la dep no
rompen (molde de los adaptadores LLM/calendar). El reporte agrega SOLO datos agregados, nunca PHI
individual (Ley N.º 29733).
"""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from typing import Any
from xml.sax.saxutils import escape

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import BadRequestException
from app.modules.dashboards.enums import ReportFormat
from app.modules.dashboards.schemas.report import ReportRequest
from app.modules.dashboards.schemas.summary import KpiSummary
from app.modules.dashboards.services import metrics as metrics_service

PDF_MEDIA_TYPE = "application/pdf"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_ALL_SECTIONS = ("kpis", "funnel", "distribution", "evolution")

# (atributo de KpiSummary, etiqueta ES, tipo de formato). tipo ∈ {pct (fracción 0–1), int, money}.
_KPI_ROWS: list[tuple[str, str, str]] = [
    ("conversion_rate", "Tasa de conversión", "pct"),
    ("lead_to_appt_rate", "Tasa lead → cita", "pct"),
    ("confirmation_rate", "Tasa de confirmación", "pct"),
    ("show_rate", "Tasa de asistencia", "pct"),
    ("no_show_rate", "Tasa de inasistencia", "pct"),
    ("customer_rate", "Conversión a cliente", "pct"),
    ("total_conversations", "Conversaciones", "int"),
    ("conversations_bot", "Conversaciones del bot", "int"),
    ("total_leads", "Leads", "int"),
    ("total_appointments", "Citas agendadas", "int"),
    ("total_confirmed", "Citas confirmadas", "int"),
    ("total_attended", "Citas atendidas", "int"),
    ("total_customers", "Clientes nuevos", "int"),
    ("bot_turns", "Turnos del bot", "int"),
    ("bot_cost_usd", "Costo estimado del bot (USD)", "money"),
]


def _wanted_sections(req: ReportRequest) -> set[str]:
    """Secciones a incluir. Filtra valores desconocidos; vacío (o todo desconocido) = TODAS."""
    wanted = {s for s in req.sections if s in _ALL_SECTIONS}
    return wanted or set(_ALL_SECTIONS)


def _filename(req: ReportRequest, ext: str) -> str:
    return f"reporte-conversion_{req.date_from}_{req.date_to}.{ext}"


def _fmt_kpi(kpi: KpiSummary, key: str, kind: str) -> str:
    """Formatea un KPI para mostrar: pct = 'XX.X %', money = '$ X.XX', int = '1 234'."""
    value = getattr(kpi, key)
    if kind == "pct":
        return f"{float(value) * 100:.1f} %"
    if kind == "money":
        return f"$ {float(value):.2f}"
    return f"{int(value):,}".replace(",", " ")


async def _resolve_branch_name(db: AsyncSession, req: ReportRequest) -> str:
    """Nombre de la sede del filtro (para la portada). None = todas; id desconocido → cae al id."""
    if not req.branch_id:
        return "Todas las sedes"
    meta = (await metrics_service.get_meta(db)).data
    for branch in meta.branches:
        if branch.id == req.branch_id:
            return branch.name
    return req.branch_id


async def _gather(db: AsyncSession, req: ReportRequest, wanted: set[str]) -> dict[str, Any]:
    """Trae los datasets pedidos del rollup (vía metrics). La 1ª llamada valida el rango
    (_guard_range dentro de cada get_*) → un rango inválido lanza DASHBOARD_INVALID_DATE_RANGE
    antes de construir el binario. `req` ES un DashboardFilter (ReportRequest lo hereda)."""
    data: dict[str, Any] = {}
    if "kpis" in wanted:
        data["kpis"] = (await metrics_service.get_summary(db, req)).data
    if "funnel" in wanted:
        data["funnel"] = (await metrics_service.get_funnel(db, req)).data
    if "distribution" in wanted:
        data["distribution"] = (await metrics_service.get_appointments_distribution(db, req)).data
    if "evolution" in wanted:
        data["evolution"] = (await metrics_service.get_leads_evolution(db, req)).data
    return data


async def generate_report(db: AsyncSession, req: ReportRequest) -> tuple[bytes, str, str]:
    """Punto de entrada: valida el formato (REPORT_FORMAT_NOT_SUPPORTED si no es pdf/excel) y
    delega al generador. Devuelve (content_bytes, media_type, filename)."""
    try:
        fmt = ReportFormat(req.format)
    except ValueError:
        raise BadRequestException(
            f"Formato de reporte no soportado: {req.format!r} (usar 'pdf' o 'excel').",
            code="REPORT_FORMAT_NOT_SUPPORTED",
        ) from None
    if fmt == ReportFormat.pdf:
        return await _generate_pdf(db, req)
    return await _generate_excel(db, req)


async def _generate_pdf(db: AsyncSession, req: ReportRequest) -> tuple[bytes, str, str]:
    # Lazy import: reportlab es dep de runtime/Docker, no del boot/smoke local.
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    wanted = _wanted_sections(req)
    data = await _gather(db, req, wanted)
    branch_name = await _resolve_branch_name(db, req)

    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    h2 = styles["Heading2"]
    body = styles["BodyText"]

    header_bg = colors.HexColor("#0F6CBD")  # azul de marca (brandPalette.primary)
    grid = colors.HexColor("#E5E9F0")
    stripe = colors.HexColor("#FAFBFD")

    def make_table(rows: list[list[str]], col_widths: list[float] | None = None) -> Table:
        table = Table(rows, colWidths=col_widths, hAlign="LEFT", repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), header_bg),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("GRID", (0, 0), (-1, -1), 0.5, grid),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, stripe]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        return table

    story: list[Any] = [
        Paragraph("Reporte de conversión", title_style),
        Paragraph("Medisage", h2),
        Spacer(1, 6 * mm),
        Paragraph(f"<b>Rango:</b> {req.date_from} a {req.date_to}", body),
        # `branch_name` (nombre de sede, dato de la clínica) se ESCAPA: el Paragraph lleva markup
        # intra-párrafo que reportlab parsea como XML → un nombre con '&'/'<'/'>' (ej. "Norte & Sur")
        # rompería el parser y daría 500. Las fechas y "Generado" son seguras (no llevan esos chars).
        Paragraph(f"<b>Sede:</b> {escape(branch_name)}", body),
        Paragraph(f"<b>Generado:</b> {datetime.now(UTC):%Y-%m-%d %H:%M} UTC", body),
        Spacer(1, 8 * mm),
    ]

    if "kpis" in data:
        story.append(Paragraph("Indicadores", h2))
        rows = [["Indicador", "Valor"]]
        rows += [[label, _fmt_kpi(data["kpis"], key, kind)] for key, label, kind in _KPI_ROWS]
        story.append(make_table(rows, [110 * mm, 50 * mm]))
        story.append(Spacer(1, 6 * mm))

    if "funnel" in data:
        story.append(Paragraph("Embudo de conversión", h2))
        rows = [["Etapa", "Conteo", "Tasa vs etapa previa"]]
        for stage in data["funnel"].stages:
            rate = f"{stage.rate_from_prev * 100:.0f} %" if stage.rate_from_prev is not None else "—"
            rows.append([stage.label, f"{stage.count:,}".replace(",", " "), rate])
        story.append(make_table(rows, [90 * mm, 35 * mm, 45 * mm]))
        story.append(Spacer(1, 6 * mm))

    if "distribution" in data:
        story.append(Paragraph("Distribución de citas", h2))
        rows = [["Estado", "Conteo"]]
        rows += [
            [b.label, f"{b.count:,}".replace(",", " ")] for b in data["distribution"].buckets
        ]
        rows.append(["Total", f"{data['distribution'].total:,}".replace(",", " ")])
        story.append(make_table(rows, [110 * mm, 50 * mm]))
        story.append(Spacer(1, 6 * mm))

    if "evolution" in data:
        story.append(Paragraph("Evolución de leads (por día)", h2))
        series = data["evolution"].series
        points = data["evolution"].points
        if series and points:
            rows = [["Fecha"] + [s.label for s in series]]
            for point in points:
                rows.append([point.date] + [str(point.values.get(s.key, 0)) for s in series])
            story.append(make_table(rows))
        else:
            story.append(Paragraph("No hay datos en el rango seleccionado.", body))

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="Reporte de conversión — Medisage",
    )
    doc.build(story)
    return buf.getvalue(), PDF_MEDIA_TYPE, _filename(req, "pdf")


async def _generate_excel(db: AsyncSession, req: ReportRequest) -> tuple[bytes, str, str]:
    # Lazy import: openpyxl es dep de runtime/Docker, no del boot/smoke local.
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wanted = _wanted_sections(req)
    data = await _gather(db, req, wanted)
    branch_name = await _resolve_branch_name(db, req)
    bold = Font(bold=True)

    wb = Workbook()
    default_sheet = wb.active
    if default_sheet is not None:
        wb.remove(default_sheet)

    # Hoja Resumen (portada): siempre presente → el workbook nunca queda sin hojas.
    info = wb.create_sheet("Resumen")
    info.append(["Reporte de conversión — Medisage"])
    info["A1"].font = bold
    info.append(["Rango", f"{req.date_from} a {req.date_to}"])
    info.append(["Sede", branch_name])
    info.append(["Generado", f"{datetime.now(UTC):%Y-%m-%d %H:%M} UTC"])

    # Las celdas llevan NÚMEROS crudos (no strings) para que el usuario re-grafique. Las tasas van
    # como fracción 0–1; el costo del bot (str en el wire) se parsea a float.
    if "kpis" in data:
        ws = wb.create_sheet("KPIs")
        ws.append(["Indicador", "Valor"])
        for cell in ws[1]:
            cell.font = bold
        for key, label, kind in _KPI_ROWS:
            value = getattr(data["kpis"], key)
            cell_value: float | int = (
                float(value) if kind in ("pct", "money") else int(value)
            )
            ws.append([label, cell_value])

    if "funnel" in data:
        ws = wb.create_sheet("Embudo")
        ws.append(["Etapa", "Conteo", "Tasa vs etapa previa"])
        for cell in ws[1]:
            cell.font = bold
        for stage in data["funnel"].stages:
            ws.append([stage.label, stage.count, stage.rate_from_prev])

    if "distribution" in data:
        ws = wb.create_sheet("Citas por estado")
        ws.append(["Estado", "Conteo"])
        for cell in ws[1]:
            cell.font = bold
        for bucket in data["distribution"].buckets:
            ws.append([bucket.label, bucket.count])
        ws.append(["Total", data["distribution"].total])

    if "evolution" in data:
        ws = wb.create_sheet("Leads por día")
        series = data["evolution"].series
        ws.append(["Fecha"] + [s.label for s in series])
        for cell in ws[1]:
            cell.font = bold
        for point in data["evolution"].points:
            ws.append([point.date] + [point.values.get(s.key, 0) for s in series])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), XLSX_MEDIA_TYPE, _filename(req, "xlsx")
